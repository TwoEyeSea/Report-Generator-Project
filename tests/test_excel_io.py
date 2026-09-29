from decimal import Decimal

import openpyxl
import pytest

from tp_drawer.excel_io import read_project, write_template
from tp_drawer.models import DEFAULT_TERMINATION_REASON

from .conftest import DEFAULT_HEADER, make_workbook


def test_read_project_basic(workbook_path):
    make_workbook(
        workbook_path,
        DEFAULT_HEADER,
        pit_numbers=[1, 2],
        layers=[
            ("Topsoil\n(Layer thickness)", ["3'-5\"", "3'-10\""]),
            ("Peat\n(Layer thickness)", ["1'-10\"", "1'-3\""]),
            ("Beach Sand\n(Layer thickness)", ["2'-1\"", "1'-9\""]),
        ],
        groundwater=["7'-2\"", "4'-11\""],
    )

    project = read_project(workbook_path)

    assert project.info.project_number == "AMR 15414"
    assert project.info.project_name == "Stanley Residence"
    assert project.info.equipment_provider == "SLIM LTD."
    assert len(project.pits) == 2

    pit1 = project.pits[0]
    assert pit1.number == "1"
    assert [layer.name for layer in pit1.layers] == ["Topsoil", "Peat", "Beach Sand"]
    assert pit1.layers[0].thickness_in == Decimal(41)  # 3'-5"
    assert pit1.total_depth_in == Decimal(41 + 22 + 25)
    assert pit1.groundwater_depth_in == Decimal(86)  # 7'-2"
    assert pit1.termination_reason == DEFAULT_TERMINATION_REASON


def test_read_project_pit_missing_a_layer(workbook_path):
    make_workbook(
        workbook_path,
        DEFAULT_HEADER,
        pit_numbers=[1, 2],
        layers=[
            ("Topsoil\n(Layer thickness)", ["3'-5\"", "3'-10\""]),
            ("Peat\n(Layer thickness)", ["1'-10\"", None]),  # pit 2 has no peat
            ("Beach Sand\n(Layer thickness)", ["2'-1\"", "1'-9\""]),
        ],
        groundwater=["7'-2\"", "4'-11\""],
    )

    project = read_project(workbook_path)

    pit2 = project.pits[1]
    assert [layer.name for layer in pit2.layers] == ["Topsoil", "Beach Sand"]


def test_read_project_groundwater_not_encountered(workbook_path):
    make_workbook(
        workbook_path,
        DEFAULT_HEADER,
        pit_numbers=[1, 2],
        layers=[("Topsoil\n(Layer thickness)", ["3'-5\"", "3'-10\""])],
        groundwater=[None, "-"],
    )

    project = read_project(workbook_path)

    assert project.pits[0].groundwater_depth_in is None
    assert project.pits[1].groundwater_depth_in is None


def test_read_project_malformed_thickness_names_cell(workbook_path):
    make_workbook(
        workbook_path,
        DEFAULT_HEADER,
        pit_numbers=[1],
        layers=[("Topsoil\n(Layer thickness)", ["not-a-dimension"])],
        groundwater=["7'-2\""],
    )

    with pytest.raises(ValueError, match="B[0-9]+"):
        read_project(workbook_path)


def test_read_project_final_strata(workbook_path):
    make_workbook(
        workbook_path,
        DEFAULT_HEADER,
        pit_numbers=[1, 2],
        layers=[("Topsoil\n(Layer thickness)", ["3'-5\"", "3'-10\""])],
        groundwater=["7'-2\"", "4'-11\""],
        final_strata=["Ironshore Formation Marl", None],
    )

    project = read_project(workbook_path)

    assert project.pits[0].termination_reason == "IRONSHORE FORMATION MARL"
    assert project.pits[1].termination_reason == DEFAULT_TERMINATION_REASON


def test_read_project_pit_notes(workbook_path):
    make_workbook(
        workbook_path,
        DEFAULT_HEADER,
        pit_numbers=[1, 2, 3],
        layers=[("Topsoil\n(Layer thickness)", ["3'-5\"", "3'-10\"", "2'-0\""])],
        groundwater=["7'-2\"", "4'-11\"", "5'-0\""],
        notes={1: "Boulders present within the spoils.", 2: "Trial pit walls unstable"},
    )

    project = read_project(workbook_path)

    assert project.pits[0].extra_note == "Boulders present within the spoils."
    assert project.pits[1].extra_note == "Trial pit walls unstable"
    assert project.pits[2].extra_note == ""  # no row for pit 3 -> untouched default


def test_read_project_pit_notes_tolerates_tiral_typo(workbook_path):
    # "Tiral Pit 1:" (transposed letters) recurs across this practice's
    # own real input sheets -- must be read the same as "Trial Pit 1:".
    path = workbook_path
    make_workbook(
        path,
        DEFAULT_HEADER,
        pit_numbers=[1],
        layers=[("Topsoil\n(Layer thickness)", ["3'-5\""])],
        groundwater=["7'-2\""],
    )
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    next_row = ws.max_row + 2
    ws.cell(row=next_row, column=1, value="Notes:")
    ws.cell(row=next_row + 1, column=1, value="Tiral Pit 1:")
    ws.cell(row=next_row + 1, column=2, value="Shallow PVC conduit encountered.")
    wb.save(path)

    project = read_project(path)

    assert project.pits[0].extra_note == "Shallow PVC conduit encountered."


def test_read_project_without_notes_section_defaults_blank(workbook_path):
    make_workbook(
        workbook_path,
        DEFAULT_HEADER,
        pit_numbers=[1],
        layers=[("Topsoil\n(Layer thickness)", ["3'-5\""])],
        groundwater=["7'-2\""],
    )

    project = read_project(workbook_path)

    assert project.pits[0].extra_note == ""


def test_read_project_walls_stable(workbook_path):
    make_workbook(
        workbook_path,
        DEFAULT_HEADER,
        pit_numbers=[1, 2, 3],
        layers=[("Topsoil\n(Layer thickness)", ["3'-5\"", "3'-10\"", "2'-0\""])],
        groundwater=["7'-2\"", "4'-11\"", "5'-0\""],
        walls_stable=["Yes", "No", None],
    )

    project = read_project(workbook_path)

    assert project.pits[0].walls_stable is True
    assert project.pits[1].walls_stable is False
    assert project.pits[2].walls_stable is True  # blank cell defaults to stable


def test_read_project_walls_stable_case_insensitive(workbook_path):
    make_workbook(
        workbook_path,
        DEFAULT_HEADER,
        pit_numbers=[1, 2],
        layers=[("Topsoil\n(Layer thickness)", ["3'-5\"", "3'-10\""])],
        groundwater=["7'-2\"", "4'-11\""],
        walls_stable=["no", "YES"],
    )

    project = read_project(workbook_path)

    assert project.pits[0].walls_stable is False
    assert project.pits[1].walls_stable is True


def test_read_project_without_walls_stable_row_defaults_stable(workbook_path):
    # No walls_stable= at all -- an older input sheet built before this
    # field existed should still load fine.
    make_workbook(
        workbook_path,
        DEFAULT_HEADER,
        pit_numbers=[1],
        layers=[("Topsoil\n(Layer thickness)", ["3'-5\""])],
        groundwater=["7'-2\""],
    )

    project = read_project(workbook_path)

    assert project.pits[0].walls_stable is True


def test_write_template_round_trips(tmp_path):
    path = tmp_path / "template.xlsx"
    write_template(path, num_pits=3, layer_names=("Topsoil", "Peat"))

    project = read_project(path)

    assert len(project.pits) == 3
    assert project.info.date == ""  # blank, ready for the user to fill in
