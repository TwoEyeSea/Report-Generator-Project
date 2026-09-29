from pathlib import Path

import openpyxl
import pytest


def make_workbook(
    path: Path,
    header_fields: dict[str, str],
    pit_numbers: list,
    layers: list[tuple[str, list]],
    groundwater: list,
    final_strata: list | None = None,
    walls_stable: list | None = None,
    notes: dict[int, str] | None = None,
) -> Path:
    """Build and save a workbook shaped like the trial pit input form.

    header_fields: ordered {label: value} for the header block.
    layers: ordered [(row_label, [value_per_pit])].
    groundwater / final_strata / walls_stable: [value_per_pit], aligned
    with pit_numbers.
    notes: {pit_number: note_text} -- written as the sheet's own
    "Notes:" section (one "Trial Pit N:" row per entry, in a fixed
    column, not one column per pit like the fields above).
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    row = 1
    for label, value in header_fields.items():
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
        row += 1

    pit_header_row = row
    ws.cell(row=pit_header_row, column=1, value="Trial pit")
    for i, number in enumerate(pit_numbers):
        ws.cell(row=pit_header_row, column=2 + i, value=number)

    row = pit_header_row + 1
    for label, values in layers:
        ws.cell(row=row, column=1, value=label)
        for i, value in enumerate(values):
            ws.cell(row=row, column=2 + i, value=value)
        row += 1

    ws.cell(row=row, column=1, value="Groundwater\n(depth below grade level)")
    for i, value in enumerate(groundwater):
        ws.cell(row=row, column=2 + i, value=value)
    row += 1

    if final_strata is not None:
        ws.cell(row=row, column=1, value="Final Strata")
        for i, value in enumerate(final_strata):
            ws.cell(row=row, column=2 + i, value=value)
        row += 1

    if walls_stable is not None:
        ws.cell(row=row, column=1, value="Pit walls stable?")
        for i, value in enumerate(walls_stable):
            ws.cell(row=row, column=2 + i, value=value)
        row += 1

    if notes is not None:
        row += 1  # a blank row, same as every real reference sheet has here
        ws.cell(row=row, column=1, value="Notes:")
        row += 1
        for pit_number, text in notes.items():
            ws.cell(row=row, column=1, value=f"Trial Pit {pit_number}:")
            ws.cell(row=row, column=2, value=text)
            row += 1

    wb.save(path)
    return path


@pytest.fixture
def workbook_path(tmp_path) -> Path:
    return tmp_path / "input.xlsx"


DEFAULT_HEADER = {
    "Date": "02/07/26",
    "Project Number": "AMR 15414",
    "Project Name": "Stanley Residence",
    "Street Address": "North Side",
    "Block and Parcel": "Block 33B Parcel 205",
    "Equipment Provider": "SLIM LTD.",
    "Equipment": "JOHN DEERE BACKHOE (24\" BUCKET)",
}
