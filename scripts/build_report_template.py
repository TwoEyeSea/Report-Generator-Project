"""One-time (well, one-time-per-source-document) build step that turns the
reference report `reference/Templates/Bagnall Pool and Cabana Report.docx`
into `data/report_template.docx` -- the same file, with its project-specific
text replaced by docxtpl/Jinja2 tags, ready to be rendered against any
Project's data by tp_drawer.report_io.generate_report.

WHY THIS IS A SEPARATE SCRIPT INSTEAD OF HAND-EDITING IN WORD:
The normal way to build a docxtpl template is to open the .docx in Word and
type "{{ tag }}" over the text you want to become dynamic. That works fine
for a human, but doing it by hand is also easy to get subtly wrong (typing
the tag with a font change partway through the word, autocorrect touching
the curly braces, etc.). Since we already have the exact source document,
scripting the same edits keeps them fully reproducible and re-runnable if
the source template is ever revised.

Every tag inserted here is written into a SINGLE run (see the helper
functions below), which sidesteps the one real hazard of docx templating:
Word routinely splits a sentence across many <w:r> runs (a paint-by-history
artifact of edits/spell-check), and a tag whose braces land in different
runs can fail to parse. Keeping each tag inside one run avoids that
entirely, and it's also simply the pattern docxtpl itself recommends.

Run this script whenever `data/report_template.docx` needs to be rebuilt
from a revised source document:
    python scripts/build_report_template.py
"""

from pathlib import Path

import docx
from docx.table import Table

from _docx_tagging_utils import add_list_numbering, clear_cell, clone_row, highlight_yellow, set_default_font, set_font_size, set_list_level, set_tag

SOURCE = Path(__file__).resolve().parent.parent / "reference" / "Templates" / "Bagnall Pool and Cabana Report.docx"
DEST = Path(__file__).resolve().parent.parent / "data" / "report_template.docx"
DEFAULT_FONT = "Roboto"
BODY_FONT_SIZE = 11


def build() -> None:
    d = docx.Document(SOURCE)
    body_paras = d.paragraphs

    # ---- Header: letterhead block (report date, project number, title) ----
    header = d.sections[0].header
    hp = header.paragraphs
    # hp[3]: "AMR CONSULTING ENGINEERS\tDate: 8th September 2026"
    # runs 0-2 = "AMR CONSULTING ENGINEERS" + tab (static letterhead) -- keep.
    # run 3 onward = "Date: 8th September 2026" -- becomes one tag.
    set_tag(hp[3], 3, "Date: {{ report_date }}")
    # hp[4]: "Unit 203 & 204 Cannon Place\tProject #: AMR 15483"
    # runs 0-3 = AMR's own office address (never project-specific) -- keep.
    # run 4 onward = "\tProject #: AMR 15483" -- becomes one tag.
    set_tag(hp[4], 4, "\tProject #: {{ project_number }}")
    # hp[12]: "PROJECT:\tGEOTECHNICAL REPORT, PROPOSED BAGNALL POOL AND CABANA – BLOCK 24A PARCEL 631, ..."
    # runs 0-1 = "PROJECT:" + tab -- keep. run 2 onward -- becomes one tag.
    set_tag(hp[12], 2, "{{ project_title_line }}", bold=True)

    # ---- INTRODUCTION ----
    # body_paras[3]: "\nThe Geotechnical investigation ... completed on the 27th of August 2026,
    # at the request of the client, Britt Bagnall, c/o Oasis Pools. The purpose ..."
    # run 0 is a leading line break -- leave it; run 1 onward is the sentence.
    set_tag(
        body_paras[3],
        1,
        "The Geotechnical investigation of the abovementioned property was completed on "
        "{{ site_work_completion_date }}, at the request of the client, {{ client }}. The purpose of "
        "this report is to provide a detailed outline of the subsurface conditions encountered beneath "
        "the site to aid in the structural design of the proposed {{ proposed_structures }}.",
    )
    # body_paras[5]: "The inspection was carried out by Mr. Alexander Moroney, MSc., of AMR Consulting Engineers."
    set_tag(
        body_paras[5],
        0,
        "The inspection was carried out by Mr. {{ engineer_name }}, of AMR Consulting Engineers.",
    )
    # body_paras[9]: "Two (2 No.) trial pits using a backhoe supplied by Slims Heavy Equipment Services Ltd., ..."
    # bold=False is deliberate, not incidental: this paragraph's first run
    # carries a stray bold from the source document's edit history (a
    # single bolded space survives even after a prior manual cleanup) --
    # collapsing to a single run without resetting bold would make the
    # *entire* generated sentence bold if that stray formatting ever comes
    # back, which is exactly the bug this fixes.
    set_tag(
        body_paras[9],
        0,
        "{{ num_pits_word }} ({{ num_pits }} No.) trial {{ pit_word }} using a {{ equipment_lower }} "
        "supplied by {{ equipment_provider }}, to assess the in-situ subsurface conditions. The soil "
        "descriptions below are based on median soil gradation and consistencies encountered during "
        "the investigation.",
        bold=False,
    )

    # ---- RESULTS: Trial Pit Results table ----
    table: Table = d.tables[0]
    data_row = table.rows[1]  # currently pit 1's hardcoded row
    old_pit2_row = table.rows[2]

    for_tr = clone_row(data_row, before=True)
    endfor_tr = clone_row(data_row, before=False)

    # Re-read the table's rows now that new <w:tr> siblings exist, so we get
    # fresh, correctly-indexed row wrapper objects rather than stale ones.
    rows = table.rows
    for_row = next(r for r in rows if r._tr is for_tr)
    endfor_row = next(r for r in rows if r._tr is endfor_tr)

    for cell in for_row.cells:
        clear_cell(cell)
    set_tag(for_row.cells[0].paragraphs[0], 0, "{%tr for pit in pits %}")
    for cell in endfor_row.cells:
        clear_cell(cell)
    set_tag(endfor_row.cells[0].paragraphs[0], 0, "{%tr endfor %}")

    set_tag(data_row.cells[0].paragraphs[0], 0, "{{ pit.number }}")
    set_tag(data_row.cells[1].paragraphs[0], 0, "{{ pit.groundwater }}")
    set_tag(data_row.cells[2].paragraphs[0], 0, "{{ pit.final_depth }}")
    set_tag(data_row.cells[3].paragraphs[0], 0, "{{ pit.stratum }}")

    # The old hardcoded pit-2 row is no longer needed -- the loop above
    # regenerates one row per pit in the project, however many there are.
    old_pit2_row._tr.getparent().remove(old_pit2_row._tr)

    # ---- DISCUSSION: subsurface commentary ----
    # body_paras[18]: bold "Subsurface Commentary – Trial Pits" heading-ish
    # lead-in (runs 0-5, kept as-is) followed by the actual commentary
    # sentence (run 6 onward) -- that part becomes one tag. Losing the bold
    # emphasis on "Native Bedrock" within the commentary is an accepted,
    # minor cosmetic simplification (see report_io.py's formation catalog).
    set_tag(body_paras[18], 6, "{{ formation_commentary }}", bold=False)

    # ---- Soil Descriptions heading: fix mis-leveled section number ----
    # This heading was carried over from the user's own template at list
    # level 2 (numId 44's level-2 format is "%1.%2.%3", i.e. "3.1.1"),
    # sitting directly under "3.0 DISCUSSION" (level 0) with no level-1
    # heading in between -- it should be level 1 ("%1.%2", i.e. "3.1"),
    # the same level "1.1 Site Work" sits at under "1.0 INTRODUCTION".
    set_list_level(body_paras[20], 1)

    # ---- Soil Descriptions: repeating per-stratum block ----
    # The template keeps exactly ONE stratum block (originally the Organic
    # Topsoil block, paragraphs 22-27 including its trailing blank spacer)
    # as the loop body, wraps it in {%p for %}/{%p endfor %}, and drops the
    # second hardcoded block (originally Marl Fill, paragraphs 28-32) --
    # the loop regenerates one block per unique layer name found in the
    # project's data, however many there are.
    strata_first = body_paras[22]
    strata_marl_first = body_paras[28]

    strata_first.insert_paragraph_before("{%p for stratum in strata %}")
    strata_marl_first.insert_paragraph_before("{%p endfor %}")

    set_tag(body_paras[22], 0, "Strata: ", bold=True)  # keep label
    set_tag(body_paras[22], 1, "{{ stratum.name }}")
    set_tag(body_paras[23], 0, "Trial pit No. encountered: ", bold=True)
    set_tag(body_paras[23], 1, "{{ stratum.pits_encountered }}")
    set_tag(body_paras[24], 0, "Thickness range: ", bold=True)
    set_tag(body_paras[24], 1, "{{ stratum.thickness_range }}")
    set_tag(body_paras[25], 0, "Underlying: ", bold=True)
    set_tag(body_paras[25], 1, "{{ stratum.underlying }}")
    set_tag(body_paras[26], 0, "Description:", bold=True)
    set_tag(body_paras[26], 1, " {{ stratum.description }}")

    for idx in range(28, 33):  # delete the old hardcoded Marl Fill block
        p = body_paras[idx]
        p._element.getparent().remove(p._element)

    # ---- NOTES: bold heading + a repeating dash-bulleted line per note ----
    # body_paras[38]: bold "All TPs:" lead-in -- the actual label text
    # ("All Tps:" or "Notes:") is data-driven, so it's now a tag.
    set_tag(body_paras[38], 0, "{{ notes_heading }}", bold=True)
    # body_paras[39]: "Excavation carried out successfully" -> the loop
    # body for however many note lines report_io.py produces (could be
    # the single default line, one combined line, or one per pit).
    notes_line_first = body_paras[39]
    notes_line_first.insert_paragraph_before("{%p for line in notes_lines %}")
    body_paras[40].insert_paragraph_before("{%p endfor %}")
    set_tag(notes_line_first, 0, "{{ line }}")

    # ---- SITE CLASSIFICATION FOR SEISMIC DESIGN (Section 4.0) ----
    # Static boilerplate (no tags) -- highlighted yellow by default like the
    # Conclusion section below: it states a default seismic site class in
    # the absence of survey data, which is the kind of assumption an
    # engineer should double-check/edit before the report goes out.
    highlight_yellow(body_paras[41])  # "SITE CLASSIFICATION FOR SEISMIC DESIGN" heading
    highlight_yellow(body_paras[43])  # the seismic site class paragraph

    # ---- CONCLUSION (Section 5.0) ----
    # The whole section -- heading, intro sentence, and the foundation
    # subsection(s) below -- gets a yellow highlight by default: it's the
    # most project-sensitive content in the report (formation name,
    # foundation choice, bearing assumptions) and the one most likely to
    # need the engineer's edits before the report goes out.
    highlight_yellow(body_paras[45])  # "CONCLUSION" heading

    # body_paras[47]: names the bearing formation and the proposed
    # structure(s) inline; the rest ("...consist of clean marl...") is
    # left as static engineering boilerplate for the engineer to
    # hand-check/edit, per the project's own stated workflow.
    set_tag(
        body_paras[47],
        0,
        "From the results of the ground investigation, the subsurface conditions across the site "
        "consist of clean marl. Given the structural loadings of the proposed {{ proposed_structures }} "
        "it is the opinion of this practice that the underlying {{ primary_formation }} may be relied "
        "upon as a bearing stratum. The recommended foundation is as follows:",
    )
    highlight_yellow(body_paras[47])

    # ---- Foundation section(s): 5.1, 5.2, ... one per recommended foundation ----
    # body_paras[49] used to be a single hand-typed "5.1 \tRecommended
    # Foundation: Strip Footings" heading. It's now the loop body for an
    # OUTER loop over report_io.py's foundation_sections (one recommended
    # foundation each), nested around the existing plain/list paragraph
    # loops -- so one recommendation renders exactly like today (one
    # "Recommended Foundation: X" subsection), and two or more each get
    # their own numbered "Option N - X" subsection, however many there are.
    #
    # The old hand-typed "5", ".1 ", tab, and label runs are dropped
    # entirely (set_tag from run 0) and replaced with a real Word
    # auto-number field (add_list_numbering) at the same list level
    # "1.1 Site Work" uses under numId 44 -- so "5.1"/"5.2"/... is always
    # computed by Word from the document's actual structure, the same
    # fragility class that broke "4.1 -> 5.1" last time this report was
    # restructured can't recur here.
    heading_para = body_paras[49]
    heading_para.insert_paragraph_before("{%p for section in foundation_sections %}")
    set_tag(heading_para, 0, "{{ section.heading_text }}", bold=True)
    add_list_numbering(heading_para, ilvl=1, numId=44)
    highlight_yellow(heading_para)

    # Plain paragraphs (50-52, no list numbering) -> nested loop over
    # section.plain_paragraphs, using paragraph 50 as the loop body.
    plain_first = body_paras[50]
    plain_first.insert_paragraph_before("{%p for para in section.plain_paragraphs %}")
    body_paras[51].insert_paragraph_before("{%p endfor %}")
    set_tag(plain_first, 0, "{{ para }}")
    highlight_yellow(plain_first)
    for idx in (51, 52):
        p = body_paras[idx]
        p._element.getparent().remove(p._element)

    # Numbered-list paragraphs (53-59, carry list numbering) -> nested loop
    # over section.list_paragraphs, using paragraph 53 (which already has
    # the list's own numPr under a *different* numId, 32) as the loop body
    # so the bullet numbering survives the loop.
    list_first = body_paras[53]
    list_first.insert_paragraph_before("{%p for para in section.list_paragraphs %}")
    # Two consecutive insert_paragraph_before calls on the same anchor
    # (body_paras[54]) land in call order immediately before it, so this
    # closes the inner list-paragraphs loop first, then the outer
    # foundation_sections loop, right after each other.
    body_paras[54].insert_paragraph_before("{%p endfor %}")  # closes inner list loop
    body_paras[54].insert_paragraph_before("{%p endfor %}")  # closes outer foundation_sections loop
    set_tag(list_first, 0, "{{ para }}")
    highlight_yellow(list_first)
    # The source document's numbered list actually runs 54-59, not 54-57
    # (the two range endpoints previously here) -- the last two items,
    # "Placement and compaction..." and "Appropriate monitoring...",
    # were never deleted, so they silently printed as static text after
    # the loop on every report regardless of which foundation was
    # recommended. Only went unnoticed because Strip Footings' own
    # catalog list_paragraphs happen to end with those same two
    # sentences, masking the bug -- a foundation with different (or no)
    # list_paragraphs, like Piles, exposed it immediately.
    for idx in (54, 55, 56, 57, 58, 59):
        p = body_paras[idx]
        p._element.getparent().remove(p._element)

    set_default_font(d, DEFAULT_FONT)

    # ---- Uniform body font size ----
    # Every non-header/footer paragraph and table cell gets forced to
    # BODY_FONT_SIZE, except the opening disclaimer paragraph (kept at
    # its own, deliberately smaller, existing 8pt). Several inserted
    # *values* -- the soil unit description, the stratum label values
    # (Strata/Trial pit No. encountered/Thickness range/Underlying), the
    # Trial Pit Results table's data cells, and a foundation write-up's
    # own paragraphs -- had no explicit size of their own and were
    # silently falling back to either Word's absolute default (~10pt) or
    # an unrelated style's size (Legal 2's 12pt, for the description),
    # both different from the rest of the body's already-11pt text.
    # Forcing every paragraph explicitly here, rather than only the
    # spots already known to be wrong, catches this same class of bug
    # wherever it next reappears -- e.g. exactly the paragraph a newly
    # added foundation catalog entry's write-up (such as Piles) flows
    # through.
    # Compare the underlying XML element (`_p`), not the Paragraph
    # wrapper object itself -- d.paragraphs builds a fresh list of new
    # Paragraph wrappers each time it's called, so two wrappers around
    # the very same paragraph are never `is` each other.
    disclaimer_p = next(p for p in d.paragraphs if p.text.strip().startswith("This Report must not be construed"))._p
    for p in d.paragraphs:
        if p._p is not disclaimer_p:
            set_font_size(p, BODY_FONT_SIZE)
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    set_font_size(p, BODY_FONT_SIZE)

    DEST.parent.mkdir(parents=True, exist_ok=True)
    d.save(DEST)
    print(f"Wrote {DEST}")


if __name__ == "__main__":
    build()
